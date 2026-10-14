import io,zipfile,hashlib
import pymupdf as fitz
import pytest
from openpyxl import load_workbook
from engine import *
from pdf_highlights import marked_pdf_zip,FINISHED
from internal_results import pending_workbook,pending_rows
from document_layout import parse_va_columns


def employee(name='MARIA HELENA MATOS',key='e'):
    return Person(key,name,'52998224725','1','EMPRESA','03.659.631/0001-05','1 - CONTRATO','08/2026','AUX SERVICOS GERAIS',False,False,'folha.pdf',1,'base')


def native_va():
    doc=fitz.open();doc.new_page().insert_text((30,60),'BOLETO preservado');doc.new_page().insert_text((30,60),'COMPROVANTE preservado');p=doc.new_page()
    for x,word in [(30,'Cartao'),(150,'Nome'),(330,'Centro'),(365,'de'),(385,'Custo'),(430,'Matricula'),(500,'Valor')]:p.insert_text((x,70),word,fontsize=8)
    for x,text in [(30,'6392XXXXXXXX1234'),(150,'MARIA HELENA MATOS'),(330,'TVV'),(430,'19'),(500,'0,00')]:p.insert_text((x,95),text,fontsize=8)
    data=doc.tobytes();doc.close();return data


def test_va_by_columns_excludes_unknown_cost_center():
    data=native_va();pages=extract_pdf('VA.pdf',data,'hash',ocr=False);rs=parse_records(pages,'VA')
    assert len(rs)==1;assert rs[0].name=='MARIA HELENA MATOS';assert rs[0].center_cost=='TVV';assert rs[0].amount==0
    assert rs[0].boxes and max(box[2] for box in rs[0].boxes)<330


def test_different_center_cost_does_not_need_hardcoded_list():
    data=native_va();doc=fitz.open(stream=data,filetype='pdf');p=doc[2]
    p.add_redact_annot(fitz.Rect(329,85,375,98),fill=(1,1,1));p.apply_redactions();p.insert_text((330,95),'XYZ-23',fontsize=8);data=doc.tobytes();doc.close()
    rs=parse_records(extract_pdf('VA.pdf',data,'new',ocr=False),'VA')
    assert rs[0].name=='MARIA HELENA MATOS';assert rs[0].center_cost=='XYZ-23'

@pytest.mark.parametrize('short',['MARIA H MATOS','MARIA HEL MAT','MARIA MATOS','M H MATOS'])
def test_abbreviation_suggestion_preserves_manual_review(short):
    e=employee();r=Record('r','VA',short,'','VA.pdf',1,'evidence')
    row=next(x for x in reconcile([e],[r],[],include_ponto=False)[0] if x['Benefício']=='VA')
    assert row['Status']=='Revisar';assert row['Nome documento']==short


def test_typos_never_override_cpf():
    e=employee();r=Record('r','Odonto','MARIA HELENA MATTOS','11144477735','odonto.pdf',1,'evidence')
    row=next(x for x in reconcile([e],[r],[],include_ponto=False)[0] if x['Benefício']=='Odonto');assert row['Status']=='Conflito CPF'


def finished_rows(e,rs):
    rows=reconcile([e],rs,[],include_ponto=False)[0]
    for row in rows:
        if row['Obrigatório'] and row['Status'] not in {'Confirmado','Validado manualmente'}:row['Status']='Ausência confirmada'
    return rows


def test_marked_pdf_keeps_filename_pages_order_text_and_only_name_boxes():
    data=native_va();pages=extract_pdf('nome original.pdf',data,'h',ocr=False);rs=parse_records(pages,'VA');e=employee();rows=finished_rows(e,rs)
    marked,log=marked_pdf_zip({'nome original.pdf':data},[e],rs,pages,rows)
    with zipfile.ZipFile(io.BytesIO(marked)) as z:
        assert 'nome original.pdf' in z.namelist();result=z.read('nome original.pdf')
    original=fitz.open(stream=data,filetype='pdf');revised=fitz.open(stream=result,filetype='pdf')
    assert len(original)==len(revised)==3;assert [p.get_text() for p in original]==[p.get_text() for p in revised]
    assert list(revised[0].annots() or [])==[];assert list(revised[1].annots() or [])==[]
    namepage=revised[2];anns=list(namepage.annots() or []);assert len(anns)==1;assert anns[0].rect.x1<330;assert 'MARCADO' in [x[4] for x in log]


def test_unresolved_reviews_block_final_zip():
    e=employee();data=native_va();pages=extract_pdf('VA.pdf',data,'h',ocr=False);rs=parse_records(pages,'VA');rows=reconcile([e],rs,pages,include_ponto=False)[0]
    with pytest.raises(ValueError):marked_pdf_zip({'VA.pdf':data},[e],rs,pages,rows)


def test_unrequired_union_not_marked():
    e=employee();data=native_va();pages=extract_pdf('Sindicato.pdf',data,'h',ocr=False);r=Record('r','Sindicato',e.name,'','Sindicato.pdf',3,'',boxes=[[150,86,270,98]])
    rows=finished_rows(e,[r]);blob,log=marked_pdf_zip({'Sindicato.pdf':data},[e],[r],pages,rows)
    with zipfile.ZipFile(io.BytesIO(blob)) as z:assert z.read('Sindicato.pdf')==data
    assert not log


def test_scanned_without_name_area_is_reported_not_guessed():
    e=employee();doc=fitz.open();doc.new_page();data=doc.tobytes();doc.close();p=Page('ponto.pdf','h',1,'',[],True,0)
    r=Record('r','Ponto',e.name,e.cpf,'ponto.pdf',1,'manual',True)
    row={'colaborador_id':e.id,'Benefício':'Ponto','Obrigatório':True,'Status':'Validado manualmente','registro_id':'r'}
    blob,log=marked_pdf_zip({'ponto.pdf':data},[e],[r],[p],[row]);assert log[0][4]=='SEM MARCAÇÃO'
    with zipfile.ZipFile(io.BytesIO(blob)) as z:assert z.read('ponto.pdf')==data


def test_manual_area_allows_scanned_point_highlight():
    e=employee();doc=fitz.open();doc.new_page();data=doc.tobytes();doc.close();p=Page('ponto.pdf','h',1,'',[],True,0)
    r=Record('r','Ponto',e.name,e.cpf,'ponto.pdf',1,'manual',True,boxes=[[80,80,250,95]])
    row={'colaborador_id':e.id,'Benefício':'Ponto','Obrigatório':True,'Status':'Validado manualmente','registro_id':'r'}
    blob,log=marked_pdf_zip({'ponto.pdf':data},[e],[r],[p],[row]);assert log[0][4]=='MARCADO'


def test_source_copies_preserved_with_original_basenames():
    data=native_va();files,msgs=unpack_files([('mesmo.pdf',data),('mesmo.pdf',data)],keep_duplicates=True)
    assert len(files)==2;assert files[0][0].endswith('mesmo.pdf') and files[1][0].endswith('mesmo.pdf')
    blob,log=marked_pdf_zip({name:data for name,data,_ in files},[],[],[],[])
    with zipfile.ZipFile(io.BytesIO(blob)) as z:assert len([n for n in z.namelist() if n.endswith('mesmo.pdf')])==2


def test_pending_excel_contains_only_unfinished_requirements():
    e=employee();r=Record('va','VA',e.name,'','VA.pdf',1,'zero',amount=0)
    rows=finished_rows(e,[r]);items=pending_rows(rows)
    assert all(x['Benefício']!='VA' for x in items);assert all(x['Status']=='Ausência confirmada' for x in items)
    wb=load_workbook(io.BytesIO(pending_workbook(rows)))
    assert wb.sheetnames==['Resumo','Pendências','Metadados'];assert wb['Pendências'].max_row==len(items)+1
    assert all(not ws.protection.sheet for ws in wb)


def test_rotated_page_preserved():
    data=native_va();doc=fitz.open(stream=data,filetype='pdf');doc[2].set_rotation(90);data=doc.tobytes();doc.close()
    pages=extract_pdf('VA.pdf',data,'r',ocr=False);rs=parse_records(pages,'VA')
    # Na rotação visual a tabela não mantém layout horizontal; preservar sem inventar marcas.
    blob,log=marked_pdf_zip({'VA.pdf':data},[],[],pages,[])
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        result=fitz.open(stream=z.read('VA.pdf'),filetype='pdf');assert len(result)==3 and result[2].rotation==90

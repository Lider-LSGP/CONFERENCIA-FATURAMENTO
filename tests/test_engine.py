import io,zipfile
import pymupdf as fitz
import pytest
from engine import *
from reports import workbook,client_package

CPF='52998224725'

def person(name='MARIA HELENA MATOS',cf=CPF,vt=False,union=False,job='AUX SERVICOS GERAIS',contract='13 - DORIO'):
    return Person(uid(cf,name,contract),name,cf,'1','EMPRESA','03.659.631/0001-05',contract,'08/2026',job,vt,union,'folha.pdf',1,'folha')

def record(name='MARIA HELENA MATOS',cf='',kind='VA',ocr=False,active=True,competence='08/2026'):
    return Record(uid(name,cf,kind),kind,name,cf,'doc.pdf',1,'registro',ocr,80 if ocr else 100,active,competence)

def run(p,rs,**kwargs):
    pp=[Page('doc.pdf','hash',1,'',['registro'],layout='Relação · VA')]
    rows,cs,extra=reconcile(p,rs,pp,include_ponto=False,**kwargs)
    return rows,cs,extra

def row(rows,b='VA',name=None):return next(r for r in rows if r['Benefício']==b and (not name or r['Colaborador']==name))

@pytest.mark.parametrize('raw,expected',[('05761687500','05761687500'),('5761687500','05761687500'),('850400538','00850400538'),('529.982.247-25',CPF),('123456789012','')])
def test_cpf_padding(raw,expected):assert cpf_normalize(raw)==expected

def test_cpf_validity():
    assert cpf_valid(CPF);assert not cpf_valid('11111111111');assert not cpf_valid('52998224724')

def test_cpf_priority_name_mismatch_review():
    p=person();r=record('JOAO PAULO SILVA',CPF)
    rows,_,_=run([p],[r]);assert row(rows)['Status']=='Revisar'

def test_exact_cpf():
    rows,_,_=run([person()],[record('MARIA H MATOS',CPF)]);assert row(rows)['Status']=='Confirmado'

def test_conflicting_cpf_never_auto():
    rows,_,_=run([person()],[record(cf='11144477735')]);assert row(rows)['Status']=='Conflito CPF'

def test_invalid_document_cpf_never_auto():
    rows,_,_=run([person()],[record(cf='00000000000')]);assert row(rows)['Status']=='Conflito CPF'

def test_name_normalization():
    rows,_,_=run([person('MÁRIA HELENA MATOS')],[record('MARIA HELENA MATOS')]);assert row(rows)['Status']=='Confirmado'

def test_abbreviation_review():
    rows,_,_=run([person()],[record('MARIA H MATOS')]);assert row(rows)['Status']=='Revisar'

def test_typo_suggestion_not_false_missing():
    rows,cs,_=run([person()],[record('MARIA HELENA MATTOS')]);assert row(rows)['Status']=='Revisar';assert cs[row(rows)['id']]

def test_ocr_requires_visual():
    rows,_,_=run([person()],[record(ocr=True)]);assert row(rows)['Status']=='Revisar'

def test_wrong_competence():
    rows,_,_=run([person()],[record(competence='07/2026')]);assert row(rows)['Status']=='Competência divergente'

def test_inactive_not_confirm():
    rows,_,_=run([person()],[record(active=False)]);assert row(rows)['Status']=='Não localizado'

def test_rules_conditional():
    rows,_,_=run([person(vt=False,union=False)],[])
    assert row(rows,'VT')['Status']=='Não exigido';assert row(rows,'Sindicato')['Status']=='Não exigido'
    assert row(rows,'Odonto')['Obrigatório']

def test_non_guard_only_gvbus():
    rows,_,_=run([person(vt=True)],[record(kind='VT Alvorada')]);assert row(rows,'VT')['Status']!='Confirmado'

def test_guard_either_operator():
    rows,_,_=run([person(vt=True,job='GUARDA VIDA')],[record(kind='VT Alvorada')]);assert row(rows,'VT')['Status']=='Confirmado'

def test_guard_both_operator():
    rows,_,_=run([person(vt=True,job='GUARDA VIDA')],[record(kind='VT Alvorada')],vt_mode='Ambas para guarda-vidas')
    assert row(rows,'VT Alvorada')['Status']=='Confirmado';assert row(rows,'VT GVBUS')['Status']!='Confirmado'

def test_homonyms_require_review():
    p1=person();p2=person(cf='11144477735')
    rows,_,_=run([p1,p2],[record()]);assert all(r['Status']=='Revisar' for r in rows if r['Benefício']=='VA')

def test_shared_record_collision():
    p1=person();p2=person('MARIA H MATOS',cf='11144477735');r=record();rid=uid(p2.id,'VA')
    rows,_,_=run([p1,p2],[r],decisions={rid:{'action':'confirm','record_id':r.id,'note':'visual'}})
    assert all(r['Status']=='Revisar' for r in rows if r['Benefício']=='VA')

def test_manual_decision_trace():
    p=person();r=record('MARIA H MATOS');rid=uid(p.id,'VA')
    rows,_,_=run([p],[r],decisions={rid:{'action':'confirm','record_id':r.id,'note':'validado','reviewer':'ana','at':'2026-10-13'}})
    assert row(rows)['Status']=='Validado manualmente';assert row(rows)['Revisor']=='ana'

def test_payroll_parser():
    lines=['Empresa: 4 - LIDER LIMPE Página: 1/1','CNPJ: 03.659.631/0001-05','Competência: 08/2026','Serviço: 13 - DORIO - CNPJ: 27.080.605/0015-91','Empr.: 1 MARIA HELENA MATOS Situação: Trabalhando CPF: 529.982.247-25 Adm: 01/01/2020','Cargo: 1 GUARDA VIDA C.B.O: 123456','258 VALE TRANSPORTE VL 10,00 10,00 D','221 MENSAL.SINDICAL 1,5% (Sindilimpe) 1,50 25,24 D']
    p=Page('f.pdf','h',1,'\n'.join(lines),lines)
    e=payroll([p])[0];assert e.vt and e.union;assert e.cpf==CPF;assert e.contract=='13 - DORIO'

def test_payroll_earning_not_discount():
    lines=['Empresa: 4 - LIDER LIMPE Página: 1/1','Empr.: 1 MARIA HELENA MATOS Situação: Trabalhando CPF: 529.982.247-25','258 VALE TRANSPORTE 10,00 10,00 P']
    assert not payroll([Page('f.pdf','h',1,'\n'.join(lines),lines)])[0].vt

def test_vt_continuation_page_and_wrap():
    p=Page('vt.pdf','h',2,'','06850001126373 10217 MARIA HELENA 0,00 153,00\nMATOS'.splitlines())
    rs=parse_records([p],'VT GVBUS');assert len(rs)==1;assert rs[0].name=='MARIA HELENA MATOS';assert rs[0].amount==153

def test_vt_zero_not_coverage():
    p=Page('vt.pdf','h',2,'',['06850001126373 10217 MARIA HELENA MATOS 0,00 0,00'])
    assert not parse_records([p],'VT GVBUS')[0].active

def test_va_wrapped_table_continuation():
    p=Page('va.pdf','h',2,'',['6392XXXXXXXX0284 MARIA HELENA MATOS PMVV 5588 R$628,31'])
    rs=parse_records([p],'VA');assert len(rs)==1;assert rs[0].name=='MARIA HELENA MATOS';assert rs[0].amount==628.31

def test_health_only_holder_header():
    lines=['Demonstrativo Analítico de Pré Pagamento.','Mensalidade - 08/2026','MARIA HELENA MATOS Cód. Referência na Empresa:','MARIA H MATOS 080.5957.003248.00-0 PLANO ASSISTENCIAL','JOAO HELENA MATOS 080.5957.003248.01-0 PLANO ASSISTENCIAL']
    r=parse_records([Page('s.pdf','h',1,'\n'.join(lines),lines)],'Saúde');assert len(r)==1;assert r[0].name=='MARIA HELENA MATOS'

def test_odontology_short_cpf():
    lines=['Demonstrativo Analítico CPF','9.0000004.60305400-4 MARIA HELENA MATOS 5 5761687500 ESSENCIAL T 38 Mensalidade Titular 10,50']
    r=parse_records([Page('o.pdf','h',1,'\n'.join(lines),lines)],'Odonto');assert r[0].cpf=='05761687500'

def test_financial_page_not_names():
    lines=['Comprovante de Pagamento Pix','Solicitante: MARIA HELENA MATOS','Nome do pagador: JOAO SILVA']
    p=Page('va.pdf','h',1,'\n'.join(lines),lines);assert parse_records([p],'VA')==[]

def test_insurance_excluded_across_pages():
    a=['Resumo do Faturamento Nr CPF','Exclusões do Mes','MARIA HELENA MATOS 529.982.247-25 01/01/1980 5,00']
    b=['Resumo do Faturamento Nr CPF','Segurados do mes Anterior','MARIA HELENA MATOS 529.982.247-25 01/01/1980 5,00']
    rs=parse_records([Page('i.pdf','h',1,'\n'.join(a),a),Page('i.pdf','h',2,'\n'.join(b),b)],'Seguro de vida');assert len(rs)==2;assert all(not r.active for r in rs)

def test_document_company_scope():
    r=record();r.company_cnpj='02.201.230/0001-44';rows,_,_=run([person()],[r]);assert row(rows)['Status']!='Confirmado'

def test_dedup_file_bytes():
    d=fitz.open();d.new_page();blob=d.tobytes();d.close()
    files,msgs=unpack_files([('a.pdf',blob),('b.pdf',blob)]);assert len(files)==1;assert msgs

def test_zip_budget():
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z:z.writestr('a.pdf',b'%PDF-'+b'x'*100)
    with pytest.raises(ValueError):unpack_files([('a.zip',b.getvalue())],max_total=20)

def test_workbook_and_client_privacy():
    from openpyxl import load_workbook
    p=person();rs=[record()];rows,_,_=run([p],rs)
    data=workbook(rows,[p],rs,[],meta={'Teste':'synthetic'});wb=load_workbook(io.BytesIO(data));assert 'Matriz geral' in wb.sheetnames;assert 'Auditoria documentos' in wb.sheetnames
    package=client_package(rows,'teste')
    with zipfile.ZipFile(io.BytesIO(package)) as z:
        html=z.read('relatorio_cliente.html').decode();assert CPF not in html;assert '***.' in html
        cli=load_workbook(io.BytesIO(z.read('conferencia_cliente.xlsx')))
        assert 'Auditoria documentos' not in cli.sheetnames

def test_export_formula_injection():
    from reports import safe_cell
    assert safe_cell('=HYPERLINK("bad")').startswith("'")

def test_columns_do_not_mix_provento_and_discount():
    line='258 VALE TRANSPORTE 10,00 10,00 P 998 INSS 8,00 20,00 D'
    assert discount_lines(line,r'VALE\s+TRANSPORTE')==[]
    assert discount_lines('219 DEV IDESBRE 0,01 0,01 P 258 VALE TRANSPORTE 10,00 10,00 D',r'VALE\s+TRANSPORTE')

def test_short_cpf_keeps_warning_not_company():
    lines=['Demonstrativo Analítico CPF','9.0000004.60305400-4 MARIA HELENA MATOS 5 5761687500 ESSENCIAL T 38 Mensalidade Titular 10,50']
    r=parse_records([Page('o.pdf','h',1,'\n'.join(lines),lines)],'Odonto')[0]
    assert r.company_cnpj=='';assert 'zeros' in r.warnings[0]

def test_manual_cannot_confirm_wrong_benefit():
    p=person();r=record(kind='Odonto');rid=uid(p.id,'VA')
    rows,_,_=run([p],[r],decisions={rid:{'action':'confirm','record_id':r.id}})
    assert row(rows)['Status']=='Revisar'

def test_ponto_competence_from_header():
    lines=['CONTROLE MENSAL DE FREQUENCIA','Agosto/2026','Colaborador: MARIA HELENA MATOS']
    r=parse_records([Page('p.pdf','h',1,'\n'.join(lines),lines,True)],'Ponto')[0]
    assert r.competence=='08/2026';assert r.ocr

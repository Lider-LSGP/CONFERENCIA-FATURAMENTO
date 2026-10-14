import io,zipfile,math,xml.etree.ElementTree as ET
import pytest
from openpyxl import load_workbook
from engine import *
from reports import workbook,matrix,GOOD,safe_cell
from point_tools import point_workspace,validate_point_link,manual_point_record


def e(name='MARIA HELENA MATOS',key='e1',month='08/2026'):
    return Person(key,name,'52998224725','1','LIDER','03.659.631/0001-05','20 - PMVV',month,'GUARDA VIDAS',True,False,'folha.pdf',1,'nome na folha')

def p(number=1,month='08/2026'):
    return Page('ponto.pdf','point_hash',number,'',['Colaborador: MARIA HELENA MATOS'],True,75,layout='Folha de ponto (OCR)',competence=month)

def docpage(lines,filename='beneficio.pdf'):
    return Page(filename,'hash',1,'\n'.join(lines),lines)

def matched(person,record,decisions=None):
    return reconcile([person],[record],[],include_ponto=True,decisions=decisions)[0]

def row(rows,benefit):return next(r for r in rows if r['Benefício']==benefit)


def test_passepay_reads_name_and_value_not_phone_or_email():
    lines=['Detalhes da Fatura A93725','61F8E9 MARIA HELENA MATOS +55 (27) 99999-0000 pessoa@example.com R$ 648,60']
    page=docpage(lines)
    rs=parse_records([page],'VT Alvorada')
    assert len(rs)==1;assert rs[0].name=='MARIA HELENA MATOS';assert rs[0].cpf=='';assert rs[0].amount==648.6;assert rs[0].registration=='61F8E9';assert 'Passepay' in page.layout


def test_alvorada_financial_pages_not_people():
    lines=['FATURA','Nome: Passepay','Compra de créditos para colaboradores R$ 10.013,22','Solicitante: MARIA HELENA MATOS']
    assert parse_records([docpage(lines)],'VT Alvorada')==[]


def test_passepay_zero_still_not_credit():
    rs=parse_records([docpage(['61F8E9 MARIA HELENA MATOS +55 (27) 99999-0000 pessoa@example.com R$ 0,00'])],'VT Alvorada')
    assert not rs[0].active


def test_va_zero_accepted_and_visible_without_waiver():
    rs=parse_records([docpage(['6392XXXXXXXX0284 MARIA HELENA MATOS PMVV 5588 R$0,00'])],'VA')
    assert rs[0].active and rs[0].amount==0
    rows=matched(e(),rs[0]);r=row(rows,'VA')
    assert r['Status']=='Confirmado';assert r['VA zerado'];assert r['Status'] in GOOD
    assert matrix(rows).iloc[0]['VA']=='Confirmado · VA zerado'
    assert 'saldo' in r['Observação'].lower()


def test_va_zero_does_not_bypass_identity_conflict():
    r=Record('zero','VA','MARIA HELENA MATOS','11144477735','va.pdf',1,'zero',amount=0)
    rows=matched(e(),r);assert row(rows,'VA')['Status']=='Conflito CPF';assert row(rows,'VA')['VA zerado']


def test_va_negative_not_accepted():
    rs=parse_records([docpage(['6392XXXXXXXX0284 MARIA HELENA MATOS PMVV 5588 R$ -10,00'])],'VA')
    assert rs[0].amount==-10 and not rs[0].active


def test_positive_va_is_preferred_when_zero_also_present():
    person=e();r1=Record('z','VA',person.name,'','va.pdf',1,'zero',amount=0);r2=Record('pos','VA',person.name,'','va.pdf',2,'credit',amount=100)
    rows=reconcile([person],[r1,r2],[],include_ponto=False)[0]
    assert not row(rows,'VA')['VA zerado'];assert row(rows,'VA')['Valor documento']==100


def test_month_not_inferred_from_payment_day():
    lines=['Créditos liberados em: 30/06/2026','61F8E9 MARIA HELENA MATOS +55 (27) 99999-0000 pessoa@example.com R$ 648,60']
    page=docpage(lines);rs=parse_records([page],'VT Alvorada');assert rs[0].competence==''
    assign_document_month([page],rs,'07/2026');assert rs[0].competence=='07/2026';assert page.competence=='07/2026'
    assert month_hint('16 VALE TRANSPORTE ALVORADA - 07.2026.pdf')=='07/2026'


def test_declared_month_conflict_blocks():
    r=Record('r','VA','MARIA HELENA MATOS','','va.pdf',1,'evidence',competence='08/2026')
    with pytest.raises(ValueError):assign_document_month([], [r],'07/2026')


def test_alvorada_short_name_needs_review():
    rs=parse_records([docpage(['61F8E9 MARIA H MATOS +55 (27) 99999-0000 pessoa@example.com R$ 648,60'])],'VT Alvorada')
    rs[0].competence='08/2026'
    assert row(matched(e(),rs[0]),'VT')['Status']=='Revisar'


def test_point_pending_hidden_after_binding():
    person=e();page=p();r=manual_point_record(person,page,'Conferi o cabeçalho.')
    decisions={uid(person.id,'Ponto'):{'action':'confirm','record_id':r.id,'note':'Conferido'}}
    rows=matched(person,r,decisions)
    pending,completed,pgs,bindings,_=point_workspace([person],[page],[r],rows)
    assert pending==[];assert completed==[person];assert bindings[(page.file,page.number)]==[person];assert pgs==[page]


def test_point_stays_pending_for_confirmed_absence():
    person=e();page=p();r=manual_point_record(person,page,'test')
    rows=matched(person,r,{uid(person.id,'Ponto'):{'action':'missing','note':'revisto'}})
    pending,completed,_,bindings,_=point_workspace([person],[page],[r],rows)
    assert pending==[person];assert completed==[];assert not bindings


def test_point_one_page_not_two_people():
    person=e();other=e('JOAO CARLOS SILVA','e2');page=p();r=manual_point_record(person,page,'revisto')
    rows=matched(person,r,{uid(person.id,'Ponto'):{'action':'confirm','record_id':r.id}})
    with pytest.raises(ValueError):validate_point_link(other,page,[r],rows)


def test_point_wrong_month_not_linkable():
    with pytest.raises(ValueError):validate_point_link(e(),p(month='07/2026'),[],[])


def test_point_clear_restores_pending_and_free_page():
    person=e();page=p();r=manual_point_record(person,page,'revisto')
    rows=matched(person,r,{uid(person.id,'Ponto'):{'action':'clear','note':'desfeito'}})
    pending,completed,_,bindings,_=point_workspace([person],[page],[r],rows)
    assert pending==[person];assert completed==[];assert not bindings;assert row(rows,'Ponto')['Status']=='Revisar'


def test_point_pages_without_ocr_are_still_available():
    page=Page('ponto.pdf','hash',2,'',[],error='OCR desabilitado',layout='Leitura inconclusiva',competence='08/2026')
    pending,_,pgs,_,_=point_workspace([e()],[page],[],[],{'ponto.pdf':'Ponto'})
    assert pgs==[page] and pending==[e()]


def test_pending_names_sorted_alphabetically():
    a=e('ZULEICA DOS SANTOS','z');b=e('ANA LUCIA SILVA','a')
    pending,_,_,_,_=point_workspace([a,b],[p()],[],[])
    assert [x.name for x in pending]==['ANA LUCIA SILVA','ZULEICA DOS SANTOS']


def test_xlsx_xml_valid_editable_and_no_structured_tables():
    person=e();r=Record('r','VA',person.name,'','va.pdf',1,'value',amount=0);rows=matched(person,r)
    data=workbook(rows,[person],[r],[],meta={'Unicode':'texto \x01 ruim','NaN':float('nan'),'danger':'=1+1'})
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert z.testzip() is None;assert not any(path.startswith('xl/tables/') for path in z.namelist())
        for path in z.namelist():
            if path.endswith('.xml') or path.endswith('.rels'):ET.fromstring(z.read(path))
    wb=load_workbook(io.BytesIO(data));assert all(not ws.protection.sheet for ws in wb)
    assert any('VA zerado' in str(c.value) for cells in wb['Matriz geral'] for c in cells)
    wb['Matriz geral']['A2']='EDITADO';roundtrip=io.BytesIO();wb.save(roundtrip)
    assert load_workbook(io.BytesIO(roundtrip.getvalue()))['Matriz geral']['A2'].value=='EDITADO'


def test_safe_cell_drops_illegal_xml_and_nan():
    assert safe_cell('a\x01b\ud800c')=='abc';assert safe_cell(float('nan'))=='';assert safe_cell(float('inf'))==''


def test_short_omitted_surnames_is_review_not_missing():
    person=e('PEDRO LUCAS SILVA MATOS');r=Record('short','VT Alvorada','PEDRO LUCAS','','alv.pdf',1,'nome curto',amount=200)
    result=matched(person,r)
    assert row(result,'VT')['Status']=='Revisar';assert row(result,'VT')['Nome documento']=='PEDRO LUCAS'


def test_subset_does_not_allow_different_first_name():
    assert not abbreviated_subset('JOAO CARLOS SILVA MATOS','PEDRO CARLOS')


def test_alvorada_repeated_list_does_not_double_match():
    person=e();a=Record('first','VT Alvorada',person.name,'','a.pdf',1,'duplicate',amount=100);b=Record('second','VT Alvorada',person.name,'','b.pdf',1,'duplicate',amount=100)
    rows=reconcile([person],[a,b],[],include_ponto=False)[0]
    assert len([r for r in rows if r['Benefício']=='VT'])==1;assert row(rows,'VT')['Status']=='Confirmado'
